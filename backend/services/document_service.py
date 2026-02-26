import mimetypes
import os
import re
import tempfile
from docx import Document
import fitz
import pdfplumber


class DocumentService:
    @staticmethod
    def _clean_text(text):
        text = re.sub(r"\s+", " ", text or "").strip()
        return text

    @staticmethod
    def read_docx(path):
        doc = Document(path)
        parts = []
        for p in doc.paragraphs:
            if p.text.strip():
                parts.append(p.text.strip())
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    value = cell.text.strip()
                    if value:
                        parts.append(value)
        return DocumentService._clean_text("\n".join(parts))

    @staticmethod
    def read_pdf(path):
        pages = []
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                txt = page.extract_text() or ""
                pages.append(txt)
        return DocumentService._clean_text("\n".join(pages))

    @staticmethod
    def read_pdf_fallback(path):
        doc = fitz.open(path)
        pages = [page.get_text("text") for page in doc]
        return DocumentService._clean_text("\n".join(pages))

    @staticmethod
    def extract_text(path):
        ext = os.path.splitext(path)[1].lower()
        if ext in [".docx", ".doc"]:
            return DocumentService.read_docx(path)
        if ext == ".pdf":
            text = DocumentService.read_pdf(path)
            if not text:
                text = DocumentService.read_pdf_fallback(path)
            return text
        raise ValueError(f"Unsupported file type: {ext}")

    @staticmethod
    def save_temp_file(file_storage):
        suffix = os.path.splitext(file_storage.filename or "")[1]
        fd, path = tempfile.mkstemp(suffix=suffix)
        os.close(fd)
        file_storage.save(path)
        return path

    @staticmethod
    def detect_mimetype(filename):
        return mimetypes.guess_type(filename)[0]
